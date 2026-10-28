"""Собственный сценарий: эффективность тренеров и организаций за сезон.

Федерация учитывает результаты тренера при премировании и аттестации (кейс, раздел 3). Платформа считает их
из своих данных: результат засчитывается тренеру, если на дату соревнования у него была связь со спортсменом
по дисциплине соревнования; организации — по организации спортсмена. Баллы — по тем же правилам, что и рейтинг.
"""

from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, time, timezone
from io import BytesIO
from typing import Annotated

from fastapi import APIRouter, Depends, Query
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from pydantic import BaseModel
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.application_models import Application, ApplicationStatus, Result, ResultStatus
from app.auth_models import User
from app.auth_routes import require_roles
from app.coaching_models import CoachLink, LinkStatus
from app.database import get_session
from app.models import Competition
from app.people import names
from app.profile_models import AthleteProfile
from app.qualification_models import AthleteRank, RankRequest
from app.rating import Entry, Rule, load_settings, match_rule
from app.reports import HEADER_FILL, HEADER_FONT, TITLE_FONT, XLSX, attachment


router = APIRouter(prefix="/api/v1", tags=["effectiveness"])
readers = require_roles("federation_admin", "ministry")


@dataclass
class Tally:
    name: str
    athletes: set[int] = field(default_factory=set)
    starts: int = 0
    podiums: int = 0
    wins: int = 0
    points: int = 0
    ranks: int = 0
    details: list[dict] = field(default_factory=list)


class Line(BaseModel):
    name: str
    athletes: int
    starts: int
    podiums: int
    wins: int
    points: int
    ranks: int
    points_per_athlete: float
    details: list[dict]


class Effectiveness(BaseModel):
    season: int
    coaches: list[Line]
    organizations: list[Line]
    method: str


def link_covers(link: CoachLink, day: date) -> bool:
    """Была ли связь тренер — спортсмен действующей на дату соревнования (по истории статусов)."""
    moment = datetime.combine(day, time(23, 59), tzinfo=timezone.utc)
    status = None
    for event in sorted(link.events, key=lambda item: item.created_at):
        created = event.created_at if event.created_at.tzinfo else event.created_at.replace(tzinfo=timezone.utc)
        if created > moment:
            break
        status = event.status
    return status in {LinkStatus.ACTIVE, LinkStatus.SUSPENDED}


def _lines(tallies: dict, limit_details: int = 30) -> list[Line]:
    lines = [
        Line(name=item.name, athletes=len(item.athletes), starts=item.starts, podiums=item.podiums, wins=item.wins, points=item.points, ranks=item.ranks,
             points_per_athlete=round(item.points / len(item.athletes), 1) if item.athletes else 0.0,
             details=sorted(item.details, key=lambda detail: detail["date"], reverse=True)[:limit_details])
        for item in tallies.values()
    ]
    return sorted(lines, key=lambda line: (-line.points, -line.podiums, line.name))


async def compute(db: AsyncSession, season: int) -> Effectiveness:
    settings = await load_settings(db)
    rules = [Rule(rule.title, rule.level, rule.competition_type_id, rule.first, rule.second, rule.third) for rule in settings.points]
    applications = list(await db.scalars(
        select(Application).join(Competition, Competition.id == Application.competition_id).join(Result, Result.application_id == Application.id)
        .where(Competition.protocol_published_at.is_not(None), Competition.date.between(date(season, 1, 1), date(season, 12, 31)),
               Application.status == ApplicationStatus.APPROVED)
    ))
    links = list(await db.scalars(select(CoachLink)))
    by_athlete: dict[int, list[CoachLink]] = defaultdict(list)
    for link in links:
        by_athlete[link.athlete_id].append(link)
    profiles = {item.user_id: item for item in await db.scalars(select(AthleteProfile))}
    people = await names(db, {link.coach_id for link in links} | {item.athlete_id for item in applications})

    coaches: dict[int, Tally] = {}
    organizations: dict[str, Tally] = {}
    for application in applications:
        competition, result = application.competition, application.result
        ok = result.result_status == ResultStatus.OK and result.place is not None
        rule = match_rule(Entry(application.athlete_id, competition.id, competition.title, competition.date, competition.level,
                                competition.competition_type_id, "", result.place or 99), rules) if ok else None
        points = rule.points(result.place) if rule and ok else 0
        detail = {"date": competition.date.isoformat(), "competition": competition.title, "athlete": people[application.athlete_id],
                  "place": result.place, "points": points}
        targets: list[Tally] = []
        for link in by_athlete[application.athlete_id]:
            if link.discipline_id == competition.discipline_id and link_covers(link, competition.date):
                targets.append(coaches.setdefault(link.coach_id, Tally(people[link.coach_id])))
        profile = profiles.get(application.athlete_id)
        org_name = profile.organization.name if profile and profile.organization else "Без организации"
        targets.append(organizations.setdefault(org_name, Tally(org_name)))
        for tally in targets:
            tally.athletes.add(application.athlete_id)
            tally.starts += 1
            tally.podiums += bool(ok and result.place <= 3)
            tally.wins += bool(ok and result.place == 1)
            tally.points += points
            tally.details.append(detail)

    # Разряды, присвоенные в сезоне: тренеру — по заявке, которую он подал; организации — по спортсмену.
    for rank in await db.scalars(select(AthleteRank).where(AthleteRank.assigned_on.between(date(season, 1, 1), date(season, 12, 31)))):
        request_submitter = None
        if rank.request_id:
            request = await db.get(RankRequest, rank.request_id)
            request_submitter = request.submitted_by_id if request and request.submitted_role == "coach" else None
        if request_submitter in coaches:
            coaches[request_submitter].ranks += 1
        profile = profiles.get(rank.athlete_id)
        org_name = profile.organization.name if profile and profile.organization else "Без организации"
        if org_name in organizations:
            organizations[org_name].ranks += 1

    return Effectiveness(
        season=season, coaches=_lines(coaches), organizations=_lines(organizations),
        method="Результат засчитывается тренеру, у которого на дату соревнования была связь со спортсменом по этой дисциплине "
               "(подтверждённая, в том числе приостановленная); организации — по организации спортсмена. Баллы — по правилам рейтинга "
               "(1–3 место с учётом уровня соревнования), без коэффициента стабильности. Разряды — присвоенные в сезоне: тренеру — "
               "по заявкам, которые он подал.",
    )


@router.get("/analytics/effectiveness", response_model=Effectiveness)
async def effectiveness(
    season: Annotated[int | None, Query(ge=2000, le=2100)] = None,
    _reader: User = Depends(readers),
    db: AsyncSession = Depends(get_session),
) -> Effectiveness:
    return await compute(db, season or date.today().year)


@router.get("/analytics/effectiveness.xlsx")
async def effectiveness_xlsx(
    season: Annotated[int | None, Query(ge=2000, le=2100)] = None,
    _reader: User = Depends(readers),
    db: AsyncSession = Depends(get_session),
) -> Response:
    report = await compute(db, season or date.today().year)
    workbook = Workbook()
    workbook.remove(workbook.active)
    headers = ["Место", "Название", "Спортсменов", "Стартов", "Призовых мест", "Побед", "Баллы", "Баллов на спортсмена", "Присвоено разрядов"]
    for title, lines in (("Тренеры", report.coaches), ("Организации", report.organizations)):
        sheet = workbook.create_sheet(title)
        sheet["A1"] = f"Эффективность: {title.lower()}, сезон {report.season}"
        sheet["A1"].font = TITLE_FONT
        for column, value in enumerate(headers, start=1):
            cell = sheet.cell(row=3, column=column, value=value)
            cell.fill, cell.font = HEADER_FILL, HEADER_FONT
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
        for index, line in enumerate(lines, start=1):
            values = [index, line.name, line.athletes, line.starts, line.podiums, line.wins, line.points, line.points_per_athlete, line.ranks]
            for column, value in enumerate(values, start=1):
                sheet.cell(row=3 + index, column=column, value=value)
        note_row = len(lines) + 5
        sheet.cell(row=note_row, column=1, value=report.method).alignment = Alignment(wrap_text=True, vertical="top")
        sheet.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=len(headers))
        sheet.row_dimensions[note_row].height = 75
        sheet.column_dimensions["B"].width = 46
        for letter in "ACDEFGHI":
            sheet.column_dimensions[letter].width = 14
        sheet.cell(row=1, column=1).font = Font(bold=True, size=14)
    buffer = BytesIO()
    workbook.save(buffer)
    return attachment(buffer.getvalue(), XLSX, f"Эффективность тренеров и организаций {report.season}.xlsx")
