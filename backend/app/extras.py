"""Дополнительные сценарии: справка с QR-кодом (Д10), документы о судействе (Д12),
показатели развития вида спорта (Д6) и доступ министерства (Д9)."""

import os
import secrets
from collections import defaultdict
from datetime import date, datetime
from io import BytesIO
from pathlib import Path as FilePath
from typing import Annotated

import qrcode
from docx.shared import Cm
from fastapi import APIRouter, Depends, File, Form, HTTPException, Path, Query, UploadFile
from fastapi.responses import FileResponse, Response
from pydantic import BaseModel, Field
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.application_models import Application, ApplicationStatus
from app.application_schemas import CompetitionBrief
from app.auth_models import AuthSession, Role, User
from app.auth_routes import current_auth_session, require_roles
from app.auth_security import normalize_email, utcnow
from app.database import get_session
from app.documents import DOCX, FEDERATION, LEVELS, PLANS, _bytes, _centered, _date_range, _document, _field, _signature, _table, result_text
from app.extras_models import Certificate, IndicatorDocument
from app.models import AgeGroup, CalendarPlan, Competition, CompetitionStatus, Organization
from app.people import athlete_briefs, names
from app.profile_models import AthleteProfile
from app.qualification_models import JudgeAssignment, JudgeProfile
from app.qualifications import CATEGORY_LABELS
from app.reports import _years, attachment, build_headcount, load_people


router = APIRouter(prefix="/api/v1", tags=["extras"])
admin_write = require_roles("federation_admin", csrf=True)
readers = require_roles("federation_admin", "ministry")


def app_url() -> str:
    return os.getenv("APP_URL", "http://localhost:5173").rstrip("/")


# =============================== справка с QR-кодом ===============================

class CertificateCreate(BaseModel):
    application_id: int = Field(gt=0)
    purpose: str | None = Field(default=None, max_length=255)


class CertificateRead(BaseModel):
    id: int
    number: str
    code: str
    verify_url: str
    issued_at: datetime
    purpose: str | None
    competition: CompetitionBrief
    place: int | None


class Verification(BaseModel):
    valid: bool
    number: str
    issued_at: datetime
    athlete: str
    birth_year: int | None
    competition: CompetitionBrief
    location: str
    age_group: str
    place: int | None
    participants: int
    result: str
    protocol_url: str


async def _participants(db: AsyncSession, application: Application) -> int:
    return await db.scalar(select(func.count()).select_from(Application).where(
        Application.competition_id == application.competition_id, Application.age_group_id == application.age_group_id,
        Application.status == ApplicationStatus.APPROVED,
    )) or 0


async def _certificate_read(db: AsyncSession, item: Certificate) -> CertificateRead:
    application = await db.get(Application, item.application_id)
    return CertificateRead(
        id=item.id, number=item.number, code=item.code, verify_url=f"{app_url()}/verify/{item.code}", issued_at=item.issued_at, purpose=item.purpose,
        competition=CompetitionBrief.model_validate(application.competition), place=application.result.place if application.result else None,
    )


@router.post("/me/certificates", response_model=CertificateRead, status_code=201)
async def create_certificate(payload: CertificateCreate, athlete: User = Depends(require_roles("athlete", csrf=True)), db: AsyncSession = Depends(get_session)) -> CertificateRead:
    """Спортсмен сам формирует справку об участии по результату в опубликованном протоколе."""
    application = await db.get(Application, payload.application_id)
    if (
        application is None or application.athlete_id != athlete.id or application.status != ApplicationStatus.APPROVED
        or application.result is None or application.competition.protocol_published_at is None
    ):
        raise HTTPException(status_code=422, detail="Справка выдаётся по результату в опубликованном протоколе")
    item = Certificate(number="—", code=secrets.token_urlsafe(18), athlete_id=athlete.id, application_id=application.id,
                       purpose=(payload.purpose or "").strip() or None, issued_at=utcnow())
    db.add(item)
    await db.flush()
    item.number = f"ФСП56-{item.issued_at.year}-{item.id:05d}"
    await db.commit()
    return await _certificate_read(db, item)


@router.get("/me/certificates", response_model=list[CertificateRead])
async def my_certificates(athlete: User = Depends(require_roles("athlete")), db: AsyncSession = Depends(get_session)) -> list[CertificateRead]:
    items = list(await db.scalars(select(Certificate).where(Certificate.athlete_id == athlete.id).order_by(Certificate.id.desc())))
    return [await _certificate_read(db, item) for item in items]


@router.get("/verify/{code}", response_model=Verification)
async def verify(code: Annotated[str, Path(min_length=10, max_length=64)], db: AsyncSession = Depends(get_session)) -> Verification:
    """Публичная проверка подлинности справки по QR-коду."""
    item = await db.scalar(select(Certificate).where(Certificate.code == code))
    if item is None:
        raise HTTPException(status_code=404, detail="Справка с таким кодом не выдавалась")
    application = await db.get(Application, item.application_id)
    brief = (await athlete_briefs(db, {item.athlete_id}))[item.athlete_id]
    result = application.result
    return Verification(
        valid=True, number=item.number, issued_at=item.issued_at, athlete=brief.full_name,
        birth_year=brief.birth_date.year if brief.birth_date else None,
        competition=CompetitionBrief.model_validate(application.competition), location=application.competition.location,
        age_group=application.age_group.name, place=result.place if result else None, participants=await _participants(db, application),
        result=result_text(result.result_status if result else None, result.score if result else None),
        protocol_url=f"{app_url()}/competitions/{application.competition_id}/protocol",
    )


def _qr_png(text: str) -> BytesIO:
    image = qrcode.make(text, box_size=8, border=2)
    buffer = BytesIO()
    image.save(buffer, format="PNG")
    buffer.seek(0)
    return buffer


@router.get("/certificates/{certificate_id}.docx")
async def certificate_docx(
    certificate_id: Annotated[int, Path(gt=0)],
    auth_session: AuthSession = Depends(current_auth_session),
    db: AsyncSession = Depends(get_session),
) -> Response:
    item = await db.get(Certificate, certificate_id)
    if item is None:
        raise HTTPException(status_code=404, detail="Справка не найдена")
    if not (auth_session.active_role == "federation_admin" or (auth_session.active_role == "athlete" and auth_session.user_id == item.athlete_id)):
        raise HTTPException(status_code=403, detail="Недостаточно прав")
    check = await verify(item.code, db)
    competition = check.competition

    document = _document()
    _centered(document, FEDERATION, size=11)
    _centered(document, f"СПРАВКА № {item.number}", bold=True, size=14, space_after=0)
    _centered(document, f"от {item.issued_at:%d.%m.%Y}", space_after=12)
    body = document.add_paragraph()
    body.paragraph_format.first_line_indent = Cm(1.25)
    body.add_run(
        f"Выдана {check.athlete}" + (f", {check.birth_year} г. р.," if check.birth_year else "") +
        f" в том, что он(а) принял(а) участие в соревновании «{competition.title}» ({LEVELS[competition.level.value]} соревнование, "
        f"{competition.competition_type.name.lower()}), дисциплина «{competition.discipline.name}», "
        f"{_date_range(competition.date, competition.end_date)}, {check.location}."
    )
    _field(document, "Возрастная группа", check.age_group)
    _field(document, "Результат", f"{check.place} место из {check.participants}" if check.place else check.result)
    if check.result and check.place:
        _field(document, "Баллы", check.result)
    if item.purpose:
        _field(document, "Справка выдана для предъявления", item.purpose)
    document.add_paragraph()
    note = document.add_paragraph("Подлинность справки можно проверить, отсканировав QR-код или открыв ссылку:")
    note.paragraph_format.space_after = 0
    document.add_paragraph(f"{app_url()}/verify/{item.code}")
    document.add_picture(_qr_png(f"{app_url()}/verify/{item.code}"), width=Cm(4))
    _signature(document, "Ответственный секретарь регионального отделения")
    return attachment(_bytes(document), DOCX, f"Справка {item.number}.docx")


# =============================== документы о судействе ===============================

@router.get("/competitions/{competition_id}/judges.docx")
async def judges_panel_docx(
    competition_id: Annotated[int, Path(gt=0)],
    _admin: User = Depends(require_roles("federation_admin")),
    db: AsyncSession = Depends(get_session),
) -> Response:
    """Справка о составе судейской коллегии соревнования."""
    competition = await db.get(Competition, competition_id)
    if competition is None:
        raise HTTPException(status_code=404, detail="Соревнование не найдено")
    panel = list(await db.scalars(select(JudgeAssignment).where(JudgeAssignment.competition_id == competition_id).order_by(JudgeAssignment.position_id, JudgeAssignment.id)))
    people = await names(db, {item.judge_id for item in panel})
    profiles = {item.user_id: item for item in await db.scalars(select(JudgeProfile).where(JudgeProfile.user_id.in_({item.judge_id for item in panel})))} if panel else {}

    document = _document()
    _centered(document, FEDERATION, size=11)
    _centered(document, "СПРАВКА", bold=True, size=14, space_after=0)
    _centered(document, "о составе судейской коллегии", bold=True, space_after=10)
    _field(document, "Соревнование", competition.title)
    _field(document, "Дисциплина", competition.discipline.name)
    _field(document, "Сроки и место", f"{_date_range(competition.date, competition.end_date)}, {competition.location}")
    plan = PLANS[competition.calendar_plan.value] + (f", № {competition.calendar_plan_number}" if competition.calendar_plan_number else "")
    _field(document, "Календарный план", plan)
    document.add_paragraph()
    rows = []
    for index, item in enumerate(panel, start=1):
        profile = profiles.get(item.judge_id)
        category = CATEGORY_LABELS[profile.category] if profile and profile.category else "без категории"
        rows.append([str(index), item.position.name, people[item.judge_id], category, f"{profile.category_expires_on:%d.%m.%Y}" if profile and profile.category_expires_on else "—"])
    if rows:
        _table(document, ["№", "Судейская должность", "Фамилия, имя, отчество", "Судейская категория", "Категория действует до"], rows, [0.8, 4.2, 5, 4.2, 2.8])
    else:
        document.add_paragraph("Судейская коллегия не сформирована.")
    _signature(document, "Председатель регионального отделения")
    return attachment(_bytes(document), DOCX, f"Судейская коллегия — {competition.title}.docx")


async def judge_report(db: AsyncSession, judge_id: int) -> Response:
    """Отчёт о соревнованиях, на которых работал судья."""
    if await db.get(User, judge_id) is None:
        raise HTTPException(status_code=404, detail="Судья не найден")
    name = (await names(db, {judge_id}))[judge_id]
    profile = await db.get(JudgeProfile, judge_id)
    assignments = list(await db.scalars(
        select(JudgeAssignment).join(Competition, Competition.id == JudgeAssignment.competition_id)
        .where(JudgeAssignment.judge_id == judge_id, Competition.date <= date.today()).order_by(Competition.date)
    ))
    document = _document()
    _centered(document, FEDERATION, size=11)
    _centered(document, "ОТЧЁТ", bold=True, size=14, space_after=0)
    _centered(document, "о соревнованиях, на которых работал спортивный судья", bold=True, space_after=10)
    _field(document, "Судья", name)
    if profile and profile.category:
        _field(document, "Судейская категория", f"{CATEGORY_LABELS[profile.category]}, {profile.order_ref}, действует до {profile.category_expires_on:%d.%m.%Y}")
    else:
        _field(document, "Судейская категория", "не присвоена")
    _field(document, "Всего соревнований", str(len(assignments)))
    document.add_paragraph()
    rows = [[f"{item.competition.date:%d.%m.%Y}", item.competition.title, LEVELS[item.competition.level.value], item.competition.discipline.name, item.position.name]
            for item in assignments]
    if rows:
        _table(document, ["Дата", "Соревнование", "Статус", "Дисциплина", "Должность"], rows, [2.2, 6.5, 2.6, 3.8, 3.2])
    else:
        document.add_paragraph("Проведённых соревнований с участием судьи нет.")
    document.add_paragraph().add_run(f"Сформировано {date.today():%d.%m.%Y} по данным платформы ФСП Оренбургской области.").italic = True
    return attachment(_bytes(document), DOCX, f"Отчёт о судействе — {name}.docx")


@router.get("/judges/{judge_id}/report.docx")
async def judge_report_admin(judge_id: Annotated[int, Path(gt=0)], _admin: User = Depends(require_roles("federation_admin")), db: AsyncSession = Depends(get_session)) -> Response:
    return await judge_report(db, judge_id)


@router.get("/me/judge/report.docx")
async def judge_report_self(judge: User = Depends(require_roles("judge")), db: AsyncSession = Depends(get_session)) -> Response:
    return await judge_report(db, judge.id)


# =============================== показатели развития ===============================

INDICATORS = {
    "headcount": ("Численность занимающихся", "реестр спортсменов: год начала и окончания занятий"),
    "organizations": ("Организации, ведущие подготовку", "реестр: организации, к которым привязаны спортсмены"),
    "competitions": ("Официальные соревнования в регионе", "календарь соревнований и опубликованные протоколы"),
}


class DocumentRead(BaseModel):
    id: int
    indicator: str
    year: int
    filename: str
    uploaded_at: datetime


class OrganizationLine(BaseModel):
    name: str
    org_type: str
    municipality: str | None
    athletes: int
    disciplines: list[str]


class CompetitionLine(BaseModel):
    id: int
    title: str
    date: date
    level: str
    discipline: str
    calendar_plan: CalendarPlan
    calendar_plan_number: str | None
    participants: int
    protocol_published: bool


class Indicator(BaseModel):
    key: str
    title: str
    source: str
    documents: list[DocumentRead]


class Development(BaseModel):
    year: int
    indicators: list[Indicator]
    headcount_years: list[int]
    headcount: list[int]
    organizations: list[OrganizationLine]
    competitions: list[CompetitionLine]
    competitions_in_plans: int


def _indicator_dir() -> FilePath:
    path = FilePath(os.getenv("UPLOAD_DIR", "uploads")) / "indicators"
    path.mkdir(parents=True, exist_ok=True)
    return path


@router.get("/development", response_model=Development)
async def development(
    year: Annotated[int | None, Query(ge=2000, le=2100)] = None,
    _reader: User = Depends(readers),
    db: AsyncSession = Depends(get_session),
) -> Development:
    """Сводка о развитии вида спорта в регионе: у каждого показателя — источник данных и подтверждающие документы."""
    year = year or date.today().year
    years = _years(year - 2, year)
    groups = list(await db.scalars(select(AgeGroup).order_by(AgeGroup.sort_order)))
    headcount = build_headcount(await load_people(db), years, groups)

    athletes = list(await db.scalars(select(AthleteProfile).where(AthleteProfile.organization_id.is_not(None))))
    per_org: dict[int, list[AthleteProfile]] = defaultdict(list)
    for athlete in athletes:
        if (athlete.start_year or 0) <= year and (athlete.end_year is None or athlete.end_year >= year):
            per_org[athlete.organization_id].append(athlete)
    organizations = []
    for org in await db.scalars(select(Organization).where(Organization.id.in_(per_org.keys())).order_by(Organization.name)) if per_org else []:
        members = per_org[org.id]
        disciplines = sorted({item.name for athlete in members for item in athlete.disciplines})
        organizations.append(OrganizationLine(name=org.name, org_type=org.org_type.value, municipality=org.municipality.name if org.municipality else None, athletes=len(members), disciplines=disciplines))

    competitions = []
    for competition in await db.scalars(select(Competition).where(
        Competition.date.between(date(year, 1, 1), date(year, 12, 31)), Competition.is_archived.is_(False),
        Competition.status != CompetitionStatus.DRAFT, Competition.date <= date.today(),
    ).order_by(Competition.date)):
        participants = await db.scalar(select(func.count()).select_from(Application).where(Application.competition_id == competition.id, Application.status == ApplicationStatus.APPROVED)) or 0
        competitions.append(CompetitionLine(
            id=competition.id, title=competition.title, date=competition.date, level=competition.level.value, discipline=competition.discipline.name,
            calendar_plan=competition.calendar_plan, calendar_plan_number=competition.calendar_plan_number, participants=participants,
            protocol_published=competition.protocol_published_at is not None,
        ))

    documents = list(await db.scalars(select(IndicatorDocument).where(IndicatorDocument.year == year).order_by(IndicatorDocument.id)))
    return Development(
        year=year,
        indicators=[Indicator(key=key, title=title, source=source, documents=[DocumentRead.model_validate(item, from_attributes=True) for item in documents if item.indicator == key])
                    for key, (title, source) in INDICATORS.items()],
        headcount_years=headcount.years, headcount=headcount.total, organizations=organizations, competitions=competitions,
        competitions_in_plans=sum(1 for item in competitions if item.calendar_plan != CalendarPlan.NONE),
    )


@router.post("/development/documents", response_model=DocumentRead, status_code=201)
async def upload_indicator_document(
    indicator: Annotated[str, Form()],
    year: Annotated[int, Form(ge=2000, le=2100)],
    file: Annotated[UploadFile, File()],
    admin: User = Depends(admin_write),
    db: AsyncSession = Depends(get_session),
) -> IndicatorDocument:
    if indicator not in INDICATORS:
        raise HTTPException(status_code=422, detail="Неизвестный показатель")
    original = FilePath(file.filename or "").name
    suffix = FilePath(original).suffix.lower()
    if suffix not in {".pdf", ".doc", ".docx", ".xls", ".xlsx", ".jpg", ".jpeg", ".png"}:
        raise HTTPException(status_code=422, detail="Допустимы PDF, Word, Excel и изображения")
    content = await file.read(10 * 1024 * 1024 + 1)
    if not content or len(content) > 10 * 1024 * 1024:
        raise HTTPException(status_code=422, detail="Файл пуст или больше 10 МБ")
    stored = f"{indicator}-{year}-{secrets.token_hex(8)}{suffix}"
    (_indicator_dir() / stored).write_bytes(content)
    item = IndicatorDocument(indicator=indicator, year=year, filename=original[:255], path=stored, uploaded_at=utcnow(), uploaded_by_id=admin.id)
    db.add(item)
    await db.commit()
    return item


@router.get("/development/documents/{document_id}")
async def download_indicator_document(document_id: Annotated[int, Path(gt=0)], _reader: User = Depends(readers), db: AsyncSession = Depends(get_session)) -> FileResponse:
    item = await db.get(IndicatorDocument, document_id)
    if item is None or not (_indicator_dir() / item.path).is_file():
        raise HTTPException(status_code=404, detail="Документ не найден")
    return FileResponse(_indicator_dir() / item.path, filename=item.filename)


@router.delete("/development/documents/{document_id}", status_code=204)
async def delete_indicator_document(document_id: Annotated[int, Path(gt=0)], _admin: User = Depends(admin_write), db: AsyncSession = Depends(get_session)) -> None:
    item = await db.get(IndicatorDocument, document_id)
    if item is None:
        raise HTTPException(status_code=404, detail="Документ не найден")
    (_indicator_dir() / item.path).unlink(missing_ok=True)
    await db.delete(item)
    await db.commit()


# =============================== доступ министерства ===============================

class MinistryAccess(BaseModel):
    email: str = Field(min_length=3, max_length=320)


class MinistryUser(BaseModel):
    id: int
    email: str
    name: str


@router.get("/admin/ministry-access", response_model=list[MinistryUser])
async def ministry_users(_admin: User = Depends(require_roles("federation_admin")), db: AsyncSession = Depends(get_session)) -> list[MinistryUser]:
    users = [user for user in await db.scalars(select(User)) if "ministry" in {role.name for role in user.roles}]
    people = await names(db, {user.id for user in users})
    return [MinistryUser(id=user.id, email=user.email, name=people[user.id]) for user in users]


@router.post("/admin/ministry-access", response_model=MinistryUser)
async def grant_ministry(payload: MinistryAccess, _admin: User = Depends(admin_write), db: AsyncSession = Depends(get_session)) -> MinistryUser:
    """Выдать сотруднику министерства доступ только для чтения (аккаунт должен быть зарегистрирован)."""
    user = await db.scalar(select(User).where(User.email == normalize_email(payload.email)))
    if user is None:
        raise HTTPException(status_code=404, detail="Пользователь с таким email не зарегистрирован")
    role = await db.scalar(select(Role).where(Role.name == "ministry"))
    if role is None:
        role = Role(name="ministry")
        db.add(role)
    if role not in user.roles:
        user.roles.append(role)
    await db.commit()
    return MinistryUser(id=user.id, email=user.email, name=(await names(db, {user.id}))[user.id])
