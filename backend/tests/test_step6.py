import unittest
from datetime import date, datetime, timedelta, timezone
from io import BytesIO
from types import SimpleNamespace

from docx import Document
from openpyxl import load_workbook
from sqlalchemy import select

from app.auth_models import Role, User
from app.auth_security import hash_password, login_limiter, utcnow
from app.coaching_models import LinkStatus
from app.effectiveness import link_covers
from support import PASSWORD, ApiDatabase, csrf, register


ADMIN = ("step6-admin@example.test", "Step6-Admin-Password-42")
TODAY = date.today()
COMPETITION_DATE = TODAY + timedelta(days=15)
SEASON = COMPETITION_DATE.year


async def seed_admin(session) -> None:
    admin = await session.scalar(select(Role).where(Role.name == "federation_admin"))
    session.add(User(email=ADMIN[0], password_hash=hash_password(ADMIN[1]), created_at=utcnow(), roles=[admin]))


def docx_text(content: bytes) -> str:
    document = Document(BytesIO(content))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        parts.extend(cell.text for row in table.rows for cell in row.cells)
    return "\n".join(parts)


class LinkCoverageTests(unittest.TestCase):
    def test_link_history_decides_credit(self) -> None:
        def at(day: date) -> datetime:
            return datetime(day.year, day.month, day.day, 12, tzinfo=timezone.utc)
        events = [SimpleNamespace(status=LinkStatus.PENDING, created_at=at(date(2025, 1, 10))),
                  SimpleNamespace(status=LinkStatus.ACTIVE, created_at=at(date(2025, 1, 12))),
                  SimpleNamespace(status=LinkStatus.ENDED, created_at=at(date(2025, 6, 1)))]
        link = SimpleNamespace(events=events)
        self.assertFalse(link_covers(link, date(2025, 1, 11)))  # ещё не подтверждена
        self.assertTrue(link_covers(link, date(2025, 3, 1)))
        self.assertFalse(link_covers(link, date(2025, 7, 1)))   # уже завершена


class Step6Tests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.db = ApiDatabase("step6")
        cls.client = cls.db.prepare(seed_admin)
        cls.ids: dict[str, int] = {}
        cls.build()

    @classmethod
    def tearDownClass(cls) -> None:
        cls.db.close(cls.client)

    def setUp(self) -> None:
        self.client.cookies.clear()
        login_limiter.clear()

    @classmethod
    def login(cls, email: str, password: str = PASSWORD) -> dict[str, str]:
        cls.client.cookies.clear()
        login_limiter.clear()
        response = cls.client.post("/api/v1/auth/login", json={"email": email, "password": password})
        assert response.status_code == 200, response.text
        return csrf(cls.client)

    @classmethod
    def build(cls) -> None:
        for key, last in (("winner", "Победителев"), ("other", "Другой")):
            cls.client.cookies.clear()
            register(cls.client, f"{key}@example.test", last_name=last)
            headers = csrf(cls.client)
            cls.client.put("/api/v1/me/profile", json={"last_name": last, "first_name": "Тест", "birth_date": "2000-02-02", "municipality_id": 1}, headers=headers)
            cls.client.put("/api/v1/me/athlete", json={"discipline_ids": [1], "start_year": SEASON - 1}, headers=headers)
            cls.ids[key] = cls.client.get("/api/v1/auth/me").json()["id"]
        cls.client.cookies.clear()
        register(cls.client, "coach6@example.test", role="coach", last_name="Тренерова")
        headers = csrf(cls.client)
        link = cls.client.post("/api/v1/coach/links", json={"athlete_id": cls.ids["winner"], "discipline_id": 1}, headers=headers).json()
        headers = cls.login("winner@example.test")
        cls.client.post(f"/api/v1/me/coach-links/{link['id']}/respond", json={"accept": True}, headers=headers)
        cls.client.cookies.clear()
        register(cls.client, "judge6@example.test", role="judge", last_name="Судьин")
        cls.ids["judge"] = cls.client.get("/api/v1/auth/me").json()["id"]
        cls.client.cookies.clear()
        register(cls.client, "ministry6@example.test", last_name="Министерская")
        cls.ids["ministry"] = cls.client.get("/api/v1/auth/me").json()["id"]

        headers = cls.login(*ADMIN)
        body = {
            "title": "Областной чемпионат для справок", "discipline_id": 1, "competition_type_id": 1, "level": "regional", "age_group_ids": [3],
            "date": COMPETITION_DATE.isoformat(), "end_date": COMPETITION_DATE.isoformat(), "location": "Оренбург", "description": "Тест",
            "format": "offline", "registration_start": (TODAY - timedelta(days=1)).isoformat(), "registration_end": (TODAY + timedelta(days=5)).isoformat(),
            "calendar_plan": "regional", "calendar_plan_number": "КП-1", "status": "registration_open",
        }
        cls.ids["competition"] = cls.client.post("/api/v1/competitions", json=body, headers=headers).json()["id"]
        cls.client.post(f"/api/v1/competitions/{cls.ids['competition']}/judges", json={"judge_id": cls.ids["judge"], "position_id": 1}, headers=headers)
        applications = {}
        for key in ("winner", "other"):
            headers = cls.login(f"{key}@example.test")
            applications[key] = cls.client.post(f"/api/v1/competitions/{cls.ids['competition']}/applications", headers=headers).json()["id"]
        headers = cls.login(*ADMIN)
        for application_id in applications.values():
            cls.client.post(f"/api/v1/applications/{application_id}/decision", json={"status": "approved"}, headers=headers)
        cls.client.put(f"/api/v1/competitions/{cls.ids['competition']}/results", json={"results": [
            {"application_id": applications["winner"], "score": 900}, {"application_id": applications["other"], "score": 500}]}, headers=headers)
        cls.client.post(f"/api/v1/competitions/{cls.ids['competition']}/protocol/publish", headers=headers)
        cls.ids.update({f"app_{key}": value for key, value in applications.items()})

    def test_certificate_with_qr_and_public_verification(self) -> None:
        headers = self.login("winner@example.test")
        self.assertEqual(self.client.post("/api/v1/me/certificates", json={"application_id": self.ids["app_other"]}, headers=headers).status_code, 422)
        created = self.client.post("/api/v1/me/certificates", json={"application_id": self.ids["app_winner"], "purpose": "для поступления в вуз"}, headers=headers)
        self.assertEqual(created.status_code, 201, created.text)
        certificate = created.json()
        self.assertRegex(certificate["number"], rf"^ФСП56-\d{{4}}-\d{{5}}$")
        self.assertTrue(certificate["verify_url"].endswith(f"/verify/{certificate['code']}"))
        self.assertEqual(len(self.client.get("/api/v1/me/certificates").json()), 1)

        document = self.client.get(f"/api/v1/certificates/{certificate['id']}.docx")
        self.assertEqual(document.status_code, 200)
        text = docx_text(document.content)
        self.assertIn(certificate["number"], text)
        self.assertIn("1 место из 2", text)
        self.assertIn("для поступления в вуз", text)
        self.assertEqual(len(Document(BytesIO(document.content)).inline_shapes), 1)  # QR-код

        self.client.cookies.clear()
        check = self.client.get(f"/api/v1/verify/{certificate['code']}").json()
        self.assertTrue(check["valid"])
        self.assertEqual((check["athlete"], check["place"], check["participants"]), ("Победителев Тест", 1, 2))
        self.assertEqual(self.client.get("/api/v1/verify/not-a-real-code-123").status_code, 404)
        self.login("other@example.test")
        self.assertEqual(self.client.get(f"/api/v1/certificates/{certificate['id']}.docx").status_code, 403)

    def test_judge_documents(self) -> None:
        self.login(*ADMIN)
        panel = docx_text(self.client.get(f"/api/v1/competitions/{self.ids['competition']}/judges.docx").content)
        self.assertIn("о составе судейской коллегии", panel)
        self.assertIn("Главный судья", panel)
        self.assertIn("Судьин Тест", panel)
        self.assertEqual(self.client.get(f"/api/v1/judges/{self.ids['judge']}/report.docx").status_code, 200)
        self.login("judge6@example.test")
        report = docx_text(self.client.get("/api/v1/me/judge/report.docx").content)
        self.assertIn("о соревнованиях, на которых работал спортивный судья", report)
        self.assertEqual(self.client.get(f"/api/v1/competitions/{self.ids['competition']}/judges.docx").status_code, 403)

    def test_ministry_reads_but_cannot_write(self) -> None:
        self.login("winner@example.test")
        self.assertEqual(self.client.get("/api/v1/development").status_code, 403)
        headers = self.login(*ADMIN)
        self.assertEqual(self.client.post("/api/v1/admin/ministry-access", json={"email": "missing@example.test"}, headers=headers).status_code, 404)
        granted = self.client.post("/api/v1/admin/ministry-access", json={"email": "MINISTRY6@example.test"}, headers=headers)
        self.assertEqual(granted.json()["id"], self.ids["ministry"])
        self.assertIn(self.ids["ministry"], [item["id"] for item in self.client.get("/api/v1/admin/ministry-access").json()])

        headers = self.login("ministry6@example.test")
        self.assertEqual(self.client.get("/api/v1/auth/me").json()["active_role"], "ministry")  # выбирается при входе
        for path in ("/api/v1/development", f"/api/v1/reports/headcount?from_year={SEASON - 1}&to_year={SEASON}", "/api/v1/teams", "/api/v1/analytics/effectiveness"):
            with self.subTest(path=path):
                self.assertEqual(self.client.get(path).status_code, 200)
        for path in ("/api/v1/rating?discipline_id=1", "/api/v1/registry/athletes", "/api/v1/audit", "/api/v1/admin/dashboard"):
            with self.subTest(path=path):
                self.assertEqual(self.client.get(path).status_code, 403)
        self.assertEqual(self.client.post("/api/v1/competitions/1/archive", headers=headers).status_code, 403)
        self.assertEqual(self.client.post("/api/v1/development/documents", data={"indicator": "headcount", "year": SEASON}, files={"file": ("a.pdf", b"%PDF", "application/pdf")}, headers=headers).status_code, 403)

    def test_development_indicators_and_documents(self) -> None:
        headers = self.login(*ADMIN)
        data = self.client.get("/api/v1/development", params={"year": SEASON}).json()
        self.assertEqual([item["key"] for item in data["indicators"]], ["headcount", "organizations", "competitions"])
        self.assertTrue(all(item["source"] for item in data["indicators"]))
        self.assertEqual(len(data["headcount"]), 3)

        wrong = self.client.post("/api/v1/development/documents", data={"indicator": "unknown", "year": SEASON}, files={"file": ("a.pdf", b"%PDF", "application/pdf")}, headers=headers)
        self.assertEqual(wrong.status_code, 422)
        uploaded = self.client.post("/api/v1/development/documents", data={"indicator": "headcount", "year": SEASON},
                                    files={"file": ("Форма 5-ФК.pdf", b"%PDF-1.4 demo", "application/pdf")}, headers=headers)
        self.assertEqual(uploaded.status_code, 201, uploaded.text)
        data = self.client.get("/api/v1/development", params={"year": SEASON}).json()
        self.assertEqual(data["indicators"][0]["documents"][0]["filename"], "Форма 5-ФК.pdf")
        self.assertEqual(self.client.get(f"/api/v1/development/documents/{uploaded.json()['id']}").content, b"%PDF-1.4 demo")
        self.assertEqual(self.client.delete(f"/api/v1/development/documents/{uploaded.json()['id']}", headers=headers).status_code, 204)

    def test_effectiveness_credits_linked_coach(self) -> None:
        self.login(*ADMIN)
        report = self.client.get("/api/v1/analytics/effectiveness", params={"season": SEASON}).json()
        coach = next(line for line in report["coaches"] if line["name"] == "Тренерова Тест")
        # Связь подтверждена до соревнования: 1 место на региональном чемпионате — 40 баллов; «Другой» тренеру не засчитан.
        self.assertEqual((coach["athletes"], coach["starts"], coach["podiums"], coach["wins"], coach["points"]), (1, 1, 1, 1, 40))
        self.assertEqual(coach["details"][0]["athlete"], "Победителев Тест")
        organization = next(line for line in report["organizations"] if line["name"] == "Без организации")
        self.assertEqual((organization["athletes"], organization["starts"], organization["points"]), (2, 2, 70))
        xlsx = self.client.get("/api/v1/analytics/effectiveness.xlsx", params={"season": SEASON})
        workbook = load_workbook(BytesIO(xlsx.content))
        self.assertEqual(workbook.sheetnames, ["Тренеры", "Организации"])
        self.assertEqual(workbook["Тренеры"]["B4"].value, "Тренерова Тест")


if __name__ == "__main__":
    unittest.main()
