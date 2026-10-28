"""Документы по формам федерации (.docx): протокол, отчёт по спортсмену, список сборной, проект приказа."""

from datetime import date
from io import BytesIO
from typing import Annotated, Literal

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from fastapi import APIRouter, Depends, HTTPException, Path
from fastapi.responses import Response
from sqlalchemy.ext.asyncio import AsyncSession

from app.applications import _competition, _protocol, my_results_for
from app.auth_models import AuthSession, User
from app.auth_routes import current_auth_session, is_active_admin, optional_auth_session, require_roles
from app.coaching import active_athlete_ids, current_rank
from app.database import get_session
from app.models import CompetitionFormat
from app.people import athlete_briefs
from app.profile_models import AthleteProfile
from app.qualifications import RANK_LABELS, order_phrase
from app.rating import get_team, read_team
from app.reports import attachment


router = APIRouter(prefix="/api/v1", tags=["documents"])
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
FEDERATION = "Региональное отделение Федерации спортивного программирования Оренбургской области"
LEVELS = {"regional": "региональное", "interregional": "межрегиональное", "all_russian": "всероссийское"}
FORMATS = {CompetitionFormat.OFFLINE: "очно", CompetitionFormat.ONLINE: "онлайн", CompetitionFormat.MIXED: "смешанный"}
PLANS = {"none": "не включено", "regional": "календарный план Оренбургской области", "ekp": "ЕКП Минспорта России"}
STATUS_TEXT = {"ok": None, "dns": "не явился", "dsq": "дисквалифицирован"}


def _document(landscape: bool = False) -> Document:
    document = Document()
    section = document.sections[0]
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
    for side in ("left_margin", "right_margin"):
        setattr(section, side, Cm(2))
    section.top_margin = section.bottom_margin = Cm(1.5)
    style = document.styles["Normal"]
    style.font.name, style.font.size = "Times New Roman", Pt(12)
    return document


def _centered(document: Document, text: str, *, bold: bool = False, size: int = 12, space_after: int = 4) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(space_after)
    run = paragraph.add_run(text)
    run.bold, run.font.size = bold, Pt(size)


def _field(document: Document, label: str, value: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.add_run(f"{label}: ").bold = True
    paragraph.add_run(value)


def _table(document: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = ""
        run = cell.paragraphs[0].add_run(text)
        run.bold, run.font.size = True, Pt(10)
    for row in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = ""
            cell.paragraphs[0].add_run(text).font.size = Pt(10)
    if widths:
        for row in table.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Cm(width)


def _signature(document: Document, position: str, name: str = "") -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(14)
    paragraph.add_run(f"{position} ").bold = True
    paragraph.add_run(f"_______________ {name}".rstrip())


def _bytes(document: Document) -> bytes:
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def result_text(status, score) -> str:
    """Баллы для участника с результатом, иначе — «не явился» / «дисквалифицирован» / прочерк."""
    if status is None:
        return "нет результата"
    if status.value == "ok":
        return str(score) if score is not None else "—"
    return STATUS_TEXT[status.value]


def _date_range(start: date, end: date) -> str:
    return f"{start:%d.%m.%Y}" if start == end else f"{start:%d.%m.%Y} — {end:%d.%m.%Y}"


# ---------- протокол по образцам Приложения № 5 к Правилам вида спорта ----------

RULES_SOURCE = (
    "Форма — рекомендуемый образец протокола из Приложения № 5 к Правилам вида спорта «спортивное программирование», "
    "утверждённым приказом Минспорта России от 11.08.2023 № 584"
)
FORM_TITLES = {
    "tasks": "по программированию алгоритмическому, программированию робототехники, программированию БАС "
             "и программированию информационных систем безопасности",
    "product": "по программированию продуктовому",
}
BLANK_TASKS = 12
BLANK_TEAMS = 3
BLANK_JUDGES = 3


def initials(name: str | None) -> str:
    """«Тимофеев Борис Сергеевич» → «Б. С. Тимофеев», как в подписи протокола."""
    if not name:
        return ""
    last, *rest = name.split()
    return " ".join([*(f"{part[0]}." for part in rest), last])


def _cell(cell, text: str, *, bold: bool = False, size: int = 9, center: bool = True) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    if center:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.bold, run.font.size = bold, Pt(size)


def _grid(document: Document, rows: int, cols: int):
    table = document.add_table(rows=rows, cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    return table


def _protocol_signatures(document: Document, chief_judge: str | None, chief_secretary: str | None) -> None:
    """Подписи как в образце: должность, подпись, инициалы и фамилия, дата и место печати."""
    for position, name in (("Главный судья соревнований", chief_judge), ("Главный секретарь соревнований", chief_secretary)):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(12)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.add_run(f"{position}\t_________________\t{initials(name) or '_____________________'}")
        caption = document.add_paragraph()
        caption.paragraph_format.space_after = Pt(0)
        caption.add_run("\t\t\t\t(подпись)\t\t(инициалы, фамилия)").font.size = Pt(8)
    document.add_paragraph().add_run("«__» __________ 20__ г.").font.size = Pt(11)
    document.add_paragraph("М.П.")


def _participant(row) -> str:
    athlete = row.athlete
    parts = [athlete.full_name]
    if athlete.birth_date:
        parts.append(f"{athlete.birth_date.year} г. р.")
    parts.append(athlete.organization or "без организации")
    return ", ".join(parts)


def _outcome(row) -> str | None:
    if row.result_status is None:
        return "нет результата"
    return STATUS_TEXT[row.result_status.value]


def _tasks_table(document: Document, rows: list, task_count: int) -> None:
    """Таблица образца: № п/п, команда (состав), задачи 1…N, «Решено», «Время выполнения задач»."""
    table = _grid(document, 2 + max(len(rows), 1), task_count + 4)
    header, numbers = table.rows[0].cells, table.rows[1].cells
    for column, title in ((0, "№ п/п"), (1, "Наименование команды (состав)")):
        table.cell(0, column).merge(table.cell(1, column))
        _cell(table.cell(0, column), title, bold=True)
    tasks = table.cell(0, 2).merge(table.cell(0, task_count + 1))
    _cell(tasks, "Задачи", bold=True)
    for index in range(task_count):
        _cell(numbers[index + 2], str(index + 1), bold=True)
    for column, title in ((task_count + 2, "Решено"), (task_count + 3, "Время выполнения задач")):
        table.cell(0, column).merge(table.cell(1, column))
        _cell(header[column], title, bold=True)
    for offset, row in enumerate(rows):
        cells = table.rows[offset + 2].cells
        _cell(cells[0], f"{row.place}." if row.place else "—")
        _cell(cells[1], _participant(row), center=False)
        marks = row.task_marks or [""] * task_count
        for index in range(task_count):
            _cell(cells[index + 2], marks[index] if index < len(marks) else "")
        outcome = _outcome(row)
        solved = str(row.solved) if row.solved is not None else (str(row.score) if row.score is not None else "")
        _cell(cells[task_count + 2], outcome or solved)
        _cell(cells[task_count + 3], "" if outcome or row.penalty_minutes is None else str(row.penalty_minutes))


def _product_table(document: Document, rows: list, criteria: list[str]) -> None:
    """Таблица образца продуктового программирования: по каждой команде — судьи 1–3 и строка «ОБЩЕЕ»."""
    lines = sum(len(row.judge_scores or []) + 1 for row in rows) or 1
    table = _grid(document, 1 + lines, 4 + len(criteria))
    header = table.rows[0].cells
    titles = ["Место", "Команда (состав)", "Фамилия, имя, отчество (при наличии) судьи", *(f"{item} (10%)" for item in criteria), "ИТОГО"]
    for cell, title in zip(header, titles):
        _cell(cell, title, bold=True, size=8)
    line = 1
    for row in rows:
        judges = row.judge_scores or []
        first = line
        for number, judge in enumerate(judges, start=1):
            cells = table.rows[line].cells
            _cell(cells[2], judge.judge or f"Судья {number}", center=False)
            for index, score in enumerate(judge.scores):
                _cell(cells[3 + index], str(score))
            _cell(cells[3 + len(criteria)], str(judge.total))
            line += 1
        cells = table.rows[line].cells
        _cell(cells[2], "ОБЩЕЕ", bold=True, center=False)
        outcome = _outcome(row)
        for index in range(len(criteria)):
            _cell(cells[3 + index], str(row.overall_scores[index]) if row.overall_scores else "", bold=True)
        _cell(cells[3 + len(criteria)], outcome or (str(row.score) if row.score is not None else ""), bold=True)
        for column in (0, 1):
            merged = table.cell(first, column).merge(table.cell(line, column)) if line > first else table.cell(first, column)
            _cell(merged, (f"{row.place}" if row.place else "—") if column == 0 else _participant(row), center=column == 0)
        line += 1


def _protocol_heading(document: Document, form: str) -> None:
    _centered(document, FEDERATION, size=11)
    _centered(document, "ПРОТОКОЛ", bold=True, size=14, space_after=0)
    _centered(document, FORM_TITLES[form], size=11, space_after=8)


@router.get("/competitions/{competition_id}/protocol.docx")
async def protocol_docx(
    competition_id: Annotated[int, Path(gt=0)],
    viewer: AuthSession | None = Depends(optional_auth_session),
    db: AsyncSession = Depends(get_session),
) -> Response:
    competition = await _competition(db, competition_id)
    if competition.protocol_published_at is None and not is_active_admin(viewer):
        raise HTTPException(status_code=404, detail="Протокол ещё не опубликован")
    protocol = await _protocol(db, competition)

    document = _document(landscape=True)
    _protocol_heading(document, protocol.form)
    _centered(document, competition.title, bold=True, space_after=10)
    _field(document, "Дисциплина", competition.discipline.name)
    _field(document, "Статус и тип", f"{LEVELS[competition.level.value]}, {competition.competition_type.name.lower()}")
    _field(document, "Сроки и место проведения", f"{_date_range(competition.date, competition.end_date)}, {competition.location} ({FORMATS[competition.format]})")
    plan = PLANS[competition.calendar_plan.value] + (f", № {competition.calendar_plan_number}" if competition.calendar_plan_number else "")
    _field(document, "Календарный план", plan)
    if protocol.published_at is None:
        _field(document, "Статус документа", "черновик, протокол не опубликован")

    task_count = protocol.task_count or max((len(row.task_marks or []) for group in protocol.groups for row in group.rows), default=0)
    for group in protocol.groups:
        heading = document.add_paragraph()
        heading.paragraph_format.space_before = Pt(10)
        heading.add_run(f"Возрастная группа: {group.age_group.name}").bold = True
        if protocol.form == "product":
            _product_table(document, group.rows, protocol.criteria)
        elif task_count:
            _tasks_table(document, group.rows, task_count)
        else:
            rows = [[str(row.place) if row.place else "—", _participant(row), result_text(row.result_status, row.score),
                     "" if row.penalty_minutes is None else str(row.penalty_minutes)] for row in group.rows]
            _table(document, ["№ п/п", "Наименование команды (состав)", "Результат (баллы)", "Время выполнения задач"], rows, [2, 12, 3, 4])
    if not protocol.groups:
        document.add_paragraph("Участников с одобренными заявками нет.")
    if protocol.tasks_summary is not None:
        summary = protocol.tasks_summary
        offered = str(summary.offered) if summary.offered else "___"
        paragraph = document.add_paragraph(
            f"К решению были предложены {offered} задач, кол-во команд, решивших все задачи: {summary.solved_all}; "
            f"кол-во команд, не решивших ни одной задачи: {summary.solved_none}"
        )
        paragraph.paragraph_format.space_before = Pt(8)
    note = document.add_paragraph()
    note.add_run(RULES_SOURCE + ".").italic = True
    note.runs[0].font.size = Pt(8)
    _protocol_signatures(document, protocol.chief_judge, protocol.chief_secretary)
    return attachment(_bytes(document), DOCX, f"Протокол — {competition.title}.docx")


@router.get("/protocol-forms/{form}.docx")
async def protocol_blank(form: Literal["tasks", "product"]) -> Response:
    """Пустой бланк рекомендуемого образца протокола (Приложение № 5 к Правилам вида спорта)."""
    document = _document(landscape=True)
    _centered(document, "Рекомендуемый образец", size=11, space_after=0)
    _centered(document, "протокола " + FORM_TITLES[form], bold=True, space_after=10)
    if form == "tasks":
        table = _grid(document, 2 + BLANK_TEAMS, BLANK_TASKS + 4)
        for column, title in ((0, "N п/п"), (1, "Наименование команды (состав)")):
            table.cell(0, column).merge(table.cell(1, column))
            _cell(table.cell(0, column), title, bold=True)
        _cell(table.cell(0, 2).merge(table.cell(0, BLANK_TASKS + 1)), "Задачи", bold=True)
        for index in range(BLANK_TASKS):
            _cell(table.rows[1].cells[index + 2], str(index + 1), bold=True)
        for column, title in ((BLANK_TASKS + 2, "Решено"), (BLANK_TASKS + 3, "Время выполнения задач")):
            table.cell(0, column).merge(table.cell(1, column))
            _cell(table.cell(0, column), title, bold=True)
        for index in range(BLANK_TEAMS):
            _cell(table.rows[2 + index].cells[0], f"{index + 1}.")
        document.add_paragraph(
            "К решению были предложены ___ задач, кол-во команд, решивших все задачи: ____; "
            "кол-во команд, не решивших ни одной задачи: ____"
        ).paragraph_format.space_before = Pt(8)
    else:
        criteria = ["Работоспособность прототипа", "Соответствие функциональным требованиям", "Презентация", "Технологичность", "Потенциал"]
        table = _grid(document, 1 + BLANK_TEAMS * (BLANK_JUDGES + 1), 3 + len(criteria))
        titles = ["Фамилия, имя, отчество (при наличии) судьи", "Команда (состав)", *(f"{item} (10%)" for item in criteria), "ИТОГО"]
        for cell, title in zip(table.rows[0].cells, titles):
            _cell(cell, title, bold=True, size=8)
        line = 1
        for _team in range(BLANK_TEAMS):
            for number in range(1, BLANK_JUDGES + 1):
                _cell(table.rows[line].cells[0], f"Судья {number}", center=False)
                line += 1
            _cell(table.rows[line].cells[0], "ОБЩЕЕ", bold=True, center=False)
            line += 1
    _protocol_signatures(document, None, None)
    name = "Образец протокола — задачи" if form == "tasks" else "Образец протокола — продуктовое программирование"
    return attachment(_bytes(document), DOCX, f"{name}.docx")


# ---------- отчёт по спортсмену ----------

async def _can_see_athlete(db: AsyncSession, auth_session: AuthSession, athlete_id: int) -> bool:
    role = auth_session.active_role
    if role == "federation_admin":
        return True
    if role == "athlete":
        return auth_session.user_id == athlete_id
    if role == "coach":
        return athlete_id in await active_athlete_ids(db, auth_session.user_id)
    return False


@router.get("/athletes/{athlete_id}/report.docx")
async def athlete_report(
    athlete_id: Annotated[int, Path(gt=0)],
    auth_session: AuthSession = Depends(current_auth_session),
    db: AsyncSession = Depends(get_session),
) -> Response:
    """Отчёт о соревнованиях, в которых участвовал спортсмен: сам спортсмен, его тренер или федерация."""
    if await db.get(AthleteProfile, athlete_id) is None:
        raise HTTPException(status_code=404, detail="Спортсмен не найден")
    if not await _can_see_athlete(db, auth_session, athlete_id):
        raise HTTPException(status_code=403, detail="Недостаточно прав")
    brief = (await athlete_briefs(db, {athlete_id}))[athlete_id]
    history = await my_results_for(db, athlete_id)
    rank = await current_rank(db, athlete_id)

    document = _document()
    _centered(document, FEDERATION, size=11)
    _centered(document, "ОТЧЁТ", bold=True, size=14, space_after=0)
    _centered(document, "о соревнованиях, в которых участвовал спортсмен", bold=True, space_after=10)
    _field(document, "Спортсмен", brief.full_name)
    _field(document, "Дата рождения", f"{brief.birth_date:%d.%m.%Y}" if brief.birth_date else "не указана")
    _field(document, "Организация", brief.organization or "без организации")
    _field(document, "Муниципалитет", brief.municipality or "не указан")
    _field(document, "Спортивный разряд", f"{RANK_LABELS[rank.rank]}, {order_phrase(rank.order_ref)}, действует до {rank.expires_on:%d.%m.%Y}" if rank else "не присвоен")
    document.add_paragraph()
    rows = [
        [f"{item.competition.date:%d.%m.%Y}", item.competition.title, LEVELS[item.competition.level.value], item.competition.discipline.name,
         item.age_group.name, f"{item.place} из {item.participants}" if item.place else "—",
         result_text(item.result_status, item.score)]
        for item in history
    ]
    if rows:
        _table(document, ["Дата", "Соревнование", "Статус", "Дисциплина", "Возрастная группа", "Место", "Баллы"], rows, [2.2, 4.5, 2.4, 3, 2.6, 1.8, 1.6])
    else:
        document.add_paragraph("Результатов в опубликованных протоколах нет.")
    document.add_paragraph().add_run(f"Сформировано {date.today():%d.%m.%Y} по данным платформы ФСП Оренбургской области.").italic = True
    return attachment(_bytes(document), DOCX, f"Отчёт о выступлениях — {brief.full_name}.docx")


@router.get("/me/report.docx")
async def my_report(auth_session: AuthSession = Depends(current_auth_session), db: AsyncSession = Depends(get_session)) -> Response:
    return await athlete_report(auth_session.user_id, auth_session, db)


# ---------- сборная ----------

def _team_heading(team) -> str:
    group = f", возрастная группа «{team.age_group.name}»" if team.age_group else ""
    return f"по виду спорта «спортивное программирование», дисциплина «{team.discipline.name}»{group}, на {team.season} год"


def _team_rows(team) -> list[list[str]]:
    return [
        [str(index), member.athlete.full_name, str(member.athlete.birth_date.year) if member.athlete.birth_date else "—",
         RANK_LABELS[member.rank] if member.rank else "—", member.athlete.organization or "без организации", member.athlete.municipality or "—",
         str(member.rating_place or "—"), str(member.rating_points or "—"), "резерв" if member.is_reserve else "основной"]
        for index, member in enumerate(team.members, start=1)
    ]


TEAM_HEADERS = ["№", "Фамилия, имя, отчество", "Год рождения", "Разряд", "Организация", "Муниципалитет", "Место в рейтинге", "Баллы", "Состав"]
TEAM_WIDTHS = [0.8, 5.5, 1.8, 2.2, 5, 3.5, 1.8, 1.6, 2]


@router.get("/teams/{team_id}/list.docx")
async def team_list(team_id: Annotated[int, Path(gt=0)], _reader: User = Depends(require_roles("federation_admin", "ministry")), db: AsyncSession = Depends(get_session)) -> Response:
    team = await read_team(db, await get_team(db, team_id))
    document = _document(landscape=True)
    _centered(document, FEDERATION, size=11)
    _centered(document, "СПИСОК", bold=True, size=14, space_after=0)
    _centered(document, f"спортивной сборной команды Оренбургской области {_team_heading(team)}", bold=True, space_after=10)
    _table(document, TEAM_HEADERS, _team_rows(team), TEAM_WIDTHS)
    document.add_paragraph().add_run("Состав сформирован по рейтингу спортсменов: сумма баллов за места в протоколах сезона с учётом уровня соревнований и коэффициента стабильности.").italic = True
    _signature(document, "Председатель регионального отделения")
    return attachment(_bytes(document), DOCX, f"Список сборной — {team.title}.docx")


@router.get("/teams/{team_id}/order.docx")
async def team_order(team_id: Annotated[int, Path(gt=0)], _reader: User = Depends(require_roles("federation_admin", "ministry")), db: AsyncSession = Depends(get_session)) -> Response:
    team = await read_team(db, await get_team(db, team_id))
    document = _document()
    _centered(document, FEDERATION, bold=True, size=12)
    _centered(document, "ПРОЕКТ", size=10)
    _centered(document, "ПРИКАЗ", bold=True, size=14, space_after=0)
    paragraph = document.add_paragraph("«___» ____________ 20__ г.\t\t\t\t\t\t№ ______")
    paragraph.paragraph_format.space_after = Pt(12)
    _centered(document, f"О формировании спортивной сборной команды Оренбургской области {_team_heading(team)}", bold=True, space_after=12)
    body = document.add_paragraph(
        "В соответствии с Федеральным законом от 04.12.2007 № 329-ФЗ «О физической культуре и спорте в Российской Федерации», "
        f"на основании рейтинга спортсменов Оренбургской области за {team.season} год"
    )
    body.paragraph_format.first_line_indent = Cm(1.25)
    document.add_paragraph().add_run("ПРИКАЗЫВАЮ:").bold = True
    main = sum(1 for member in team.members if not member.is_reserve)
    reserve = len(team.members) - main
    for number, text in enumerate([
        f"Утвердить список спортивной сборной команды Оренбургской области {_team_heading(team)} "
        f"в количестве {main} основных спортсменов" + (f" и {reserve} в резерве" if reserve else "") + " (приложение).",
        "Тренерам спортсменов, включённых в состав, обеспечить их подготовку и участие в соревнованиях.",
        "Контроль за исполнением настоящего приказа оставляю за собой.",
    ], start=1):
        item = document.add_paragraph(f"{number}. {text}")
        item.paragraph_format.first_line_indent = Cm(1.25)
    _signature(document, "Председатель регионального отделения")

    document.add_page_break()
    _centered(document, "Приложение к приказу от «___» ____________ 20__ г. № ______", size=10)
    _centered(document, f"Состав спортивной сборной команды Оренбургской области {_team_heading(team)}", bold=True, space_after=8)
    _table(document, TEAM_HEADERS, _team_rows(team), [0.7, 4.3, 1.4, 1.8, 3.2, 2.4, 1.3, 1.3, 1.6])
    return attachment(_bytes(document), DOCX, f"Проект приказа — {team.title}.docx")
