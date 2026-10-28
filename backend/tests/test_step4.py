import unittest
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace

from docx import Document
from openpyxl import load_workbook
from sqlalchemy import select

from app.auth_models import Role, User
from app.auth_security import hash_password, login_limiter, utcnow
from app.models import CompetitionLevel
from app.rating import Entry, Rule, compute_rating, match_rule
from app.reports import Person, build_headcount
from support import PASSWORD, ApiDatabase, csrf, register


ADMIN = ("step4-admin@example.test", "Step4-Admin-Password-42")
TODAY = date.today()
COMPETITION_DATE = TODAY + timedelta(days=20)
SEASON = COMPETITION_DATE.year


async def seed_admin(session) -> None:
    admin = await session.scalar(select(Role).where(Role.name == "federation_admin"))
    session.add(User(email=ADMIN[0], password_hash=hash_password(ADMIN[1]), created_at=utcnow(), roles=[admin]))


def docx_text(content: bytes) -> str:
    document = Document(BytesIO(content))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        parts.extend(cell.text for row in table.rows for cell in row.cells)
    return "\n".join(parts)


RULES = [
    Rule("Чемпионат России", CompetitionLevel.ALL_RUSSIAN, 1, 100, 80, 60),
    Rule("Первенство России", CompetitionLevel.ALL_RUSSIAN, 2, 80, 60, 40),
    Rule("Межрегиональные", CompetitionLevel.INTERREGIONAL, None, 60, 45, 30),
    Rule("Региональные", CompetitionLevel.REGIONAL, None, 40, 30, 20),
]


def entry(athlete: int, competition: int, level: CompetitionLevel, place: int, type_id: int = 1) -> Entry:
    return Entry(athlete, competition, f"Соревнование {competition}", date(2026, 1, competition), level, type_id, "тип", place)


class RatingRulesTests(unittest.TestCase):
    def test_rule_matching_prefers_exact_type(self) -> None:
        self.assertEqual(match_rule(entry(1, 1, CompetitionLevel.ALL_RUSSIAN, 1, type_id=2), RULES).title, "Первенство России")
        self.assertEqual(match_rule(entry(1, 1, CompetitionLevel.REGIONAL, 1, type_id=5), RULES).title, "Региональные")
        self.assertIsNone(match_rule(entry(1, 1, CompetitionLevel.ALL_RUSSIAN, 1, type_id=3), RULES))

    def test_points_coefficient_and_shared_places(self) -> None:
        scores = compute_rating([
            entry(1, 1, CompetitionLevel.REGIONAL, 1),         # 40
            entry(1, 2, CompetitionLevel.INTERREGIONAL, 2),    # 45, два призовых → ×1,2 = 102
            entry(2, 3, CompetitionLevel.ALL_RUSSIAN, 3),      # 60 за 3 место Чемпионата России
            entry(2, 1, CompetitionLevel.REGIONAL, 4),         # 4 место — 0 баллов
            entry(3, 2, CompetitionLevel.INTERREGIONAL, 3),    # 30
            entry(4, 1, CompetitionLevel.REGIONAL, 3),         # 20
            entry(4, 4, CompetitionLevel.REGIONAL, 5),         # 0 → итог 20... и одно призовое — без коэффициента
            entry(5, 5, CompetitionLevel.REGIONAL, 2),         # 30 — делит место с №3
        ], RULES, Decimal("1.2"), 2)
        table = {score.athlete_id: (score.place, score.total, score.coefficient) for score in scores}
        self.assertEqual(table[1], (1, Decimal("102.00"), Decimal("1.2")))
        self.assertEqual(table[2], (2, Decimal("60.00"), Decimal("1")))
        self.assertEqual(table[3][0], table[5][0])
        self.assertEqual(table[3][1], Decimal("30.00"))
        self.assertEqual(table[4], (5, Decimal("20.00"), Decimal("1")))
        self.assertEqual([item.points for item in scores[0].items], [40, 45])


class HeadcountTests(unittest.TestCase):
    def test_dynamics_and_sections(self) -> None:
        people = [
            Person(2010, 2024, None, "Лицей № 1", "г. Оренбург", ["Алгоритмическое", "ИИ"], "Начальная подготовка"),
            Person(2008, 2025, None, "Лицей № 1", "г. Орск", ["Алгоритмическое"], "Не указано"),
            Person(2000, 2023, 2024, "Без организации", "г. Оренбург", ["Продуктовое"], "Не указано"),
            Person(None, 2026, None, "Клуб", "Не указано", [], "Не указано"),
        ]
        groups = [SimpleNamespace(name="10–13", min_age=10, max_age=13), SimpleNamespace(name="14–17", min_age=14, max_age=17), SimpleNamespace(name="16+", min_age=16, max_age=None)]
        report = build_headcount(people, [2024, 2025, 2026], groups)
        self.assertEqual(report.total, [2, 2, 3])
        sections = {section.key: {line.name: line.counts for line in section.rows} for section in report.sections}
        self.assertEqual(sections["organizations"]["Лицей № 1"], [1, 2, 2])
        self.assertEqual(sections["disciplines"]["Алгоритмическое"], [1, 2, 2])
        self.assertEqual(sections["disciplines"]["Продуктовое"], [1, 0, 0])
        self.assertEqual(sections["age_groups"]["14–17"], [1, 2, 1])
        self.assertEqual(sections["age_groups"]["16+"], [1, 1, 2])
        self.assertEqual(sections["age_groups"]["Не указано"], [0, 0, 1])
        self.assertEqual(sections["municipalities"]["г. Оренбург"], [2, 1, 1])


class Step4ApiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.db = ApiDatabase("step4")
        cls.client = cls.db.prepare(seed_admin)
        cls.ids: dict[str, int] = {}
        cls.build_season()

    @classmethod
    def tearDownClass(cls) -> None:
        cls.db.close(cls.client)

    def setUp(self) -> None:
        self.client.cookies.clear()
        login_limiter.clear()

    @classmethod
    def login(cls, email: str, password: str = PASSWORD) -> dict[str, str]:
        cls.client.cookies.clear()
        login_limiter.clear()
        assert cls.client.post("/api/v1/auth/login", json={"email": email, "password": password}).status_code == 200
        return csrf(cls.client)

    @classmethod
    def athlete(cls, key: str, last_name: str, birth_year: int) -> None:
        cls.client.cookies.clear()
        register(cls.client, f"{key}@example.test", last_name=last_name)
        headers = csrf(cls.client)
        cls.client.put("/api/v1/me/profile", json={"last_name": last_name, "first_name": "Тест", "birth_date": f"{birth_year}-03-03", "municipality_id": 1}, headers=headers)
        body = {"discipline_ids": [1], "start_year": SEASON - 1}
        if TODAY.year - birth_year < 19:
            body |= {"guardian_name": "Представитель", "guardian_phone": "+7 900 000-00-00", "guardian_consent": True}
        assert cls.client.put("/api/v1/me/athlete", json=body, headers=headers).status_code == 200
        cls.ids[key] = cls.client.get("/api/v1/auth/me").json()["id"]

    @classmethod
    def competition(cls, title: str, level: str, places: dict[str, int]) -> int:
        headers = cls.login(*ADMIN)
        body = {
            "title": title, "discipline_id": 1, "competition_type_id": 1, "level": level, "age_group_ids": [3],
            "date": COMPETITION_DATE.isoformat(), "end_date": COMPETITION_DATE.isoformat(), "location": "Оренбург", "description": "Тест",
            "format": "offline", "registration_start": (TODAY - timedelta(days=1)).isoformat(), "registration_end": (TODAY + timedelta(days=5)).isoformat(),
            "calendar_plan": "ekp" if level != "regional" else "regional", "calendar_plan_number": "№ 1", "status": "registration_open",
        }
        competition_id = cls.client.post("/api/v1/competitions", json=body, headers=headers).json()["id"]
        applications = {}
        for key in places:
            headers = cls.login(f"{key}@example.test")
            applications[key] = cls.client.post(f"/api/v1/competitions/{competition_id}/applications", headers=headers).json()["id"]
        headers = cls.login(*ADMIN)
        for application_id in applications.values():
            cls.client.post(f"/api/v1/applications/{application_id}/decision", json={"status": "approved"}, headers=headers)
        results = [{"application_id": applications[key], "score": 1000 - place} for key, place in places.items()]
        assert cls.client.put(f"/api/v1/competitions/{competition_id}/results", json={"results": results}, headers=headers).status_code == 200
        assert cls.client.post(f"/api/v1/competitions/{competition_id}/protocol/publish", headers=headers).status_code == 200
        return competition_id

    @classmethod
    def build_season(cls) -> None:
        cls.athlete("leader", "Лидеров", SEASON - 20)
        cls.athlete("second", "Вторая", SEASON - 17)
        cls.athlete("banned", "Отстранённый", SEASON - 19)
        cls.athlete("young", "Юный", SEASON - 16)
        cls.ids["regional"] = cls.competition("Чемпионат области", "regional", {"leader": 1, "second": 2, "banned": 3, "young": 4})
        cls.ids["interregional"] = cls.competition("Межрегиональный кубок", "interregional", {"second": 1, "leader": 2})

    def test_rating_with_coefficient_filters_and_sanctions(self) -> None:
        headers = self.login(*ADMIN)
        rating = self.client.get("/api/v1/rating", params={"discipline_id": 1, "season": SEASON}).json()
        rows = {row["athlete"]["id"]: row for row in rating["rows"]}
        # Вторая: 2 место на региональном (30) + 1 место на межрегиональном (60), два призовых → ×1,2 = 108.
        self.assertEqual((rows[self.ids["second"]]["place"], rows[self.ids["second"]]["total"]), (1, "108.00"))
        leader = rows[self.ids["leader"]]
        self.assertEqual(leader["place"], 2)
        self.assertEqual((leader["base_points"], leader["coefficient"], leader["total"]), (85, "1.20", "102.00"))
        self.assertEqual([item["points"] for item in leader["items"]], [40, 45])
        self.assertEqual(rows[self.ids["young"]]["total"], "0.00")

        sanction = self.client.post("/api/v1/sanctions", json={"athlete_id": self.ids["banned"], "kind": "disqualification", "starts_on": TODAY.isoformat(), "ends_on": (TODAY + timedelta(days=90)).isoformat(), "reason": "Нарушение правил соревнований"}, headers=headers)
        self.assertEqual(sanction.status_code, 201, sanction.text)
        rating = self.client.get("/api/v1/rating", params={"discipline_id": 1, "season": SEASON}).json()
        self.assertNotIn(self.ids["banned"], [row["athlete"]["id"] for row in rating["rows"]])
        self.assertIn("Дисквалификация", rating["excluded"][0]["reason"])

        # Возрастная группа 14–17: в сезоне проходят только спортсмены 2009–2012 г. р.
        youth = self.client.get("/api/v1/rating", params={"discipline_id": 1, "season": SEASON, "age_group_id": 2}).json()
        self.assertEqual({row["athlete"]["id"] for row in youth["rows"]}, {self.ids["second"], self.ids["young"]})
        self.assertEqual(self.client.get("/api/v1/rating", params={"discipline_id": 1, "season": SEASON, "min_rank": "third"}).json()["rows"], [])

        self.assertEqual(self.client.delete(f"/api/v1/sanctions/{sanction.json()['id']}", headers=headers).status_code, 204)

    def test_rating_settings_change_points(self) -> None:
        headers = self.login(*ADMIN)
        settings = self.client.get("/api/v1/rating/settings").json()
        self.addCleanup(lambda: self.client.put("/api/v1/rating/settings", json=settings, headers=self.login(*ADMIN)))
        self.assertEqual(settings["stability_coefficient"], "1.20")
        self.assertEqual(len(settings["points"]), 4)
        changed = settings | {"stability_coefficient": "1.5", "points": [rule | ({"first": 50} if rule["level"] == "regional" else {}) for rule in settings["points"]]}
        self.assertEqual(self.client.put("/api/v1/rating/settings", json=changed, headers=headers).status_code, 200)
        leader = next(row for row in self.client.get("/api/v1/rating", params={"discipline_id": 1, "season": SEASON}).json()["rows"] if row["athlete"]["id"] == self.ids["leader"])
        self.assertEqual(leader["total"], "142.50")  # (50 + 45) × 1,5
        duplicate = settings | {"points": settings["points"] + [settings["points"][-1]]}
        self.assertEqual(self.client.put("/api/v1/rating/settings", json=duplicate, headers=headers).status_code, 422)

    def test_team_list_and_order(self) -> None:
        headers = self.login(*ADMIN)
        body = {"title": "Сборная по алгоритмическому программированию", "discipline_id": 1, "season": SEASON,
                "members": [{"athlete_id": self.ids["second"]}, {"athlete_id": self.ids["leader"]}, {"athlete_id": self.ids["young"], "is_reserve": True}]}
        team = self.client.post("/api/v1/teams", json=body, headers=headers)
        self.assertEqual(team.status_code, 201, team.text)
        members = team.json()["members"]
        self.assertEqual([member["athlete"]["id"] for member in members], [self.ids["second"], self.ids["leader"], self.ids["young"]])
        self.assertEqual(members[0]["rating_points"], "108.00")
        self.assertTrue(members[2]["is_reserve"])
        self.assertEqual(self.client.post("/api/v1/teams", json=body | {"members": [{"athlete_id": 999}]}, headers=headers).status_code, 422)

        team_id = team.json()["id"]
        listing = self.client.get(f"/api/v1/teams/{team_id}/list.docx")
        self.assertEqual(listing.status_code, 200)
        self.assertIn("filename*=UTF-8''", listing.headers["content-disposition"])
        text = docx_text(listing.content)
        self.assertIn("Лидеров Тест", text)
        self.assertIn("резерв", text)
        order = docx_text(self.client.get(f"/api/v1/teams/{team_id}/order.docx").content)
        self.assertIn("ПРИКАЗЫВАЮ", order)
        self.assertIn("в количестве 2 основных спортсменов и 1 в резерве", order)
        self.assertEqual(len(self.client.get("/api/v1/teams").json()), 1)
        self.assertEqual(self.client.delete(f"/api/v1/teams/{team_id}", headers=headers).status_code, 204)

    def test_protocol_and_athlete_report_documents(self) -> None:
        self.client.cookies.clear()
        protocol = self.client.get(f"/api/v1/competitions/{self.ids['regional']}/protocol.docx")
        self.assertEqual(protocol.status_code, 200)
        text = docx_text(protocol.content)
        for fragment in ("ПРОТОКОЛ", "Приложения № 5", "Чемпионат области", "Лидеров Тест", "999", "Главный судья соревнований", "М.П.", "календарный план"):
            self.assertIn(fragment, text)

        self.assertEqual(self.client.get(f"/api/v1/athletes/{self.ids['leader']}/report.docx").status_code, 401)
        self.login("leader@example.test")
        report = docx_text(self.client.get("/api/v1/me/report.docx").content)
        self.assertIn("Межрегиональный кубок", report)
        self.assertIn("1 из 4", report)
        self.assertIn("2 из 2", report)
        self.assertEqual(self.client.get(f"/api/v1/athletes/{self.ids['second']}/report.docx").status_code, 403)
        self.login(*ADMIN)
        self.assertEqual(self.client.get(f"/api/v1/athletes/{self.ids['second']}/report.docx").status_code, 200)

    def test_registry_and_headcount_xlsx(self) -> None:
        headers = self.login(*ADMIN)
        registry = self.client.get("/api/v1/registry/athletes").json()
        self.assertGreaterEqual(len(registry), 4)
        self.assertEqual([row["full_name"] for row in self.client.get("/api/v1/registry/athletes", params={"q": "Лидер"}).json()], ["Лидеров Тест"])
        self.assertEqual(len(self.client.get("/api/v1/registry/athletes", params={"organization_id": 0}).json()), len(registry))
        self.assertEqual(self.client.get("/api/v1/registry/athletes", params={"discipline_id": 6}).json(), [])

        report = self.client.get("/api/v1/reports/headcount", params={"from_year": SEASON - 1, "to_year": SEASON}).json()
        self.assertEqual(report["years"], [SEASON - 1, SEASON])
        self.assertEqual(report["total"][-1], len(registry))
        self.assertEqual(self.client.get("/api/v1/reports/headcount", params={"from_year": SEASON, "to_year": SEASON}).status_code, 422)

        xlsx = self.client.get("/api/v1/reports/headcount.xlsx", params={"from_year": SEASON - 1, "to_year": SEASON})
        self.assertEqual(xlsx.status_code, 200)
        workbook = load_workbook(BytesIO(xlsx.content))
        self.assertEqual(workbook.sheetnames[:3], ["Сводка", "Организации", "Дисциплины"])
        summary = workbook["Сводка"]
        self.assertIn("Численность занимающихся", summary["A1"].value)
        self.assertEqual(summary["A5"].value, "Всего занимающихся")
        self.assertEqual(summary.cell(row=5, column=3).value, len(registry))

        self.login("leader@example.test")
        for path in ("/api/v1/rating?discipline_id=1", "/api/v1/registry/athletes", "/api/v1/reports/headcount.xlsx", "/api/v1/teams"):
            with self.subTest(path=path):
                self.assertEqual(self.client.get(path).status_code, 403)
        self.assertIsNotNone(headers)


if __name__ == "__main__":
    unittest.main()
