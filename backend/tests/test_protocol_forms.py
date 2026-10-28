"""Протоколы по рекомендуемым образцам Приложения № 5 к Правилам вида спорта (приказ Минспорта № 584)."""

import unittest
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO

from docx import Document
from sqlalchemy import select

from app.application_models import Application, ApplicationStatus, ResultStatus
from app.auth_models import Role, User
from app.auth_security import hash_password, login_limiter, utcnow
from app.models import AgeGroup, CalendarPlan, Competition, CompetitionFormat, CompetitionLevel, CompetitionStatus, CompetitionType, Discipline
from app.profile_models import AthleteProfile, Profile
from app.protocol import Entry, judge_totals, normalize_mark, rank, solved_count
from support import ApiDatabase, csrf


ADMIN = ("forms-admin@example.test", "Forms-Admin-Password-42")
DAY = date.today() - timedelta(days=5)
ATHLETES = ["Задачин", "Штрафов", "Быстров", "Неявкин"]
TEAMS = ["Продуктов", "Прототипов"]


def docx_text(content: bytes) -> str:
    document = Document(BytesIO(content))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


async def seed(session) -> None:
    roles = {item.name: item for item in await session.scalars(select(Role))}
    admin = User(email=ADMIN[0], password_hash=hash_password(ADMIN[1]), created_at=utcnow(), roles=[roles["federation_admin"]])
    session.add(admin)
    disciplines = {item.code: item for item in await session.scalars(select(Discipline))}
    group = await session.scalar(select(AgeGroup).where(AgeGroup.min_age == 16))
    championship = await session.scalar(select(CompetitionType).where(CompetitionType.name == "Чемпионат"))

    def competition(title: str, code: str, task_count: int | None) -> Competition:
        item = Competition(
            title=title, discipline=disciplines[code], competition_type=championship, level=CompetitionLevel.REGIONAL, age_groups=[group],
            date=DAY, end_date=DAY, location="г. Оренбург", description="Тест", format=CompetitionFormat.OFFLINE,
            registration_start=DAY - timedelta(days=30), registration_end=DAY - timedelta(days=7), calendar_plan=CalendarPlan.REGIONAL,
            calendar_plan_number="КП-1", status=CompetitionStatus.ONGOING, is_archived=False, task_count=task_count,
        )
        session.add(item)
        return item

    algorithmic = competition("Чемпионат по задачам", "algorithmic", 4)
    product = competition("Чемпионат по продуктовому", "product", None)
    await session.flush()
    for surname, target in [*((name, algorithmic) for name in ATHLETES), *((name, product) for name in TEAMS)]:
        user = User(email=f"{surname.lower()}@example.test", password_hash=hash_password("Correct-Horse-42"), created_at=utcnow(), roles=[roles["athlete"]])
        session.add(user)
        await session.flush()
        session.add(Profile(user_id=user.id, last_name=surname, first_name="Тест", birth_date=date(2000, 1, 1), updated_at=utcnow()))
        session.add(AthleteProfile(user_id=user.id, disciplines=[target.discipline], updated_at=utcnow()))
        session.add(Application(
            competition_id=target.id, athlete_id=user.id, age_group_id=group.id, submitted_by_id=user.id, submitted_role="athlete",
            status=ApplicationStatus.APPROVED, created_at=utcnow(), decided_at=utcnow(), decided_by_id=admin.id,
        ))


class FormRulesTests(unittest.TestCase):
    def test_marks_are_normalized_and_counted(self) -> None:
        self.assertEqual([normalize_mark(mark) for mark in ["+", " +2 ", "−1", "", "+0", "-0"]], ["+", "+2", "-1", "", "+", ""])
        self.assertEqual(solved_count(["+", "+3", "-1", ""]), 2)
        with self.assertRaises(ValueError):
            normalize_mark("да")

    def test_equal_solved_tasks_are_ranked_by_time(self) -> None:
        places = rank([
            Entry(1, Decimal("3"), ResultStatus.OK, 200),
            Entry(2, Decimal("3"), ResultStatus.OK, 150),
            Entry(3, Decimal("2"), ResultStatus.OK, 10),
            Entry(4, Decimal("3"), ResultStatus.OK, 150),
        ])
        self.assertEqual(places, {2: 1, 4: 1, 1: 3, 3: 4})

    def test_product_overall_is_judges_average(self) -> None:
        judges, overall, total = judge_totals([
            {"judge": "Судья 1", "scores": ["10", "8", "6", "4", "2"]},
            {"judge": "Судья 2", "scores": ["8", "8", "8", "8", "8"]},
        ])
        self.assertEqual([item[2] for item in judges], [Decimal("30"), Decimal("40")])
        self.assertEqual(overall, [Decimal("9.00"), Decimal("8.00"), Decimal("7.00"), Decimal("6.00"), Decimal("5.00")])
        self.assertEqual(total, Decimal("35.00"))


class ProtocolFormsApiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.db = ApiDatabase("protocol_forms")
        cls.client = cls.db.prepare(seed)

    @classmethod
    def tearDownClass(cls) -> None:
        cls.db.close(cls.client)

    def setUp(self) -> None:
        login_limiter.clear()
        self.client.cookies.clear()
        self.assertEqual(self.client.post("/api/v1/auth/login", json={"email": ADMIN[0], "password": ADMIN[1]}).status_code, 200)
        self.headers = csrf(self.client)
        competitions = {item["title"]: item for item in self.client.get("/api/v1/competitions").json()}
        self.tasks_competition = competitions["Чемпионат по задачам"]
        self.product_competition = competitions["Чемпионат по продуктовому"]

    def applications(self, competition_id: int) -> dict[str, int]:
        items = self.client.get(f"/api/v1/competitions/{competition_id}/applications").json()
        return {item["athlete"]["full_name"].split()[0]: item["id"] for item in items}

    def test_tasks_protocol(self) -> None:
        competition_id = self.tasks_competition["id"]
        self.assertEqual(self.tasks_competition["protocol_form"], "tasks")
        ids = self.applications(competition_id)
        url = f"/api/v1/competitions/{competition_id}/results"
        wrong = self.client.put(url, json={"results": [{"application_id": ids["Задачин"], "task_marks": ["+", "+"]}]}, headers=self.headers)
        self.assertEqual(wrong.status_code, 422)
        self.assertIn("всем 4 задачам", wrong.text)
        judges = self.client.put(url, json={"results": [{"application_id": ids["Задачин"], "judge_scores": [{"judge": "А", "scores": [1, 1, 1, 1, 1]}]}]}, headers=self.headers)
        self.assertEqual(judges.status_code, 422)

        results = [
            {"application_id": ids["Задачин"], "task_marks": ["+", "+1", "+", "+"], "penalty_minutes": 240},
            {"application_id": ids["Штрафов"], "task_marks": ["+", "-2", "+3", ""], "penalty_minutes": 300},
            {"application_id": ids["Быстров"], "task_marks": ["+", "+", "", "-1"], "penalty_minutes": 90},
            {"application_id": ids["Неявкин"], "result_status": "dns"},
        ]
        protocol = self.client.put(url, json={"results": results}, headers=self.headers)
        self.assertEqual(protocol.status_code, 200, protocol.text)
        body = protocol.json()
        self.assertEqual(body["form"], "tasks")
        rows = {row["athlete"]["full_name"].split()[0]: row for row in body["groups"][0]["rows"]}
        self.assertEqual((rows["Задачин"]["place"], rows["Задачин"]["solved"]), (1, 4))
        # Равное число решённых задач: выше тот, у кого меньше время выполнения.
        self.assertEqual((rows["Быстров"]["place"], rows["Штрафов"]["place"]), (2, 3))
        self.assertEqual(rows["Неявкин"]["place"], None)
        self.assertEqual(body["tasks_summary"], {"offered": 4, "solved_all": 1, "solved_none": 0})

        document = docx_text(self.client.get(f"/api/v1/competitions/{competition_id}/protocol.docx").content)
        for fragment in ("ПРОТОКОЛ", "по программированию алгоритмическому", "Задачи", "Решено", "Время выполнения задач",
                         "К решению были предложены 4 задач", "решивших все задачи: 1", "+1", "не явился", "Главный секретарь соревнований", "М.П."):
            self.assertIn(fragment, document)

        locked = self.client.patch(f"/api/v1/competitions/{competition_id}", json={"task_count": 5}, headers=self.headers)
        self.assertEqual(locked.status_code, 409)

    def test_product_protocol(self) -> None:
        competition_id = self.product_competition["id"]
        self.assertEqual(self.product_competition["protocol_form"], "product")
        ids = self.applications(competition_id)
        url = f"/api/v1/competitions/{competition_id}/results"
        tasks = self.client.put(url, json={"results": [{"application_id": ids["Продуктов"], "task_marks": ["+"]}]}, headers=self.headers)
        self.assertEqual(tasks.status_code, 422)
        results = [
            {"application_id": ids["Продуктов"], "judge_scores": [
                {"judge": "Тимофеев Б. С.", "scores": [9, 8, 7, 9, 8]}, {"judge": "Абрамова Е. А.", "scores": [8, 8, 8, 8, 8]},
                {"judge": "Калинин Р. О.", "scores": [10, 9, 8, 7, 6]},
            ]},
            {"application_id": ids["Прототипов"], "judge_scores": [{"judge": "Тимофеев Б. С.", "scores": [5, 5, 5, 5, 5]}]},
        ]
        protocol = self.client.put(url, json={"results": results}, headers=self.headers)
        self.assertEqual(protocol.status_code, 200, protocol.text)
        body = protocol.json()
        self.assertEqual(body["form"], "product")
        self.assertEqual(len(body["criteria"]), 5)
        leader = body["groups"][0]["rows"][0]
        self.assertEqual((leader["place"], leader["score"]), (1, "40.33"))
        self.assertEqual([judge["total"] for judge in leader["judge_scores"]], ["41", "40", "40"])
        self.assertEqual(leader["overall_scores"][0], "9.00")
        self.assertIsNone(body["tasks_summary"])

        document = docx_text(self.client.get(f"/api/v1/competitions/{competition_id}/protocol.docx").content)
        for fragment in ("по программированию продуктовому", "Работоспособность прототипа (10%)", "ИТОГО", "ОБЩЕЕ", "Тимофеев Б. С.", "40.33"):
            self.assertIn(fragment, document)

    def test_blank_forms_are_public(self) -> None:
        self.client.cookies.clear()
        tasks = docx_text(self.client.get("/api/v1/protocol-forms/tasks.docx").content)
        for fragment in ("Рекомендуемый образец", "N п/п", "Наименование команды (состав)", "12", "Решено", "К решению были предложены ___ задач"):
            self.assertIn(fragment, tasks)
        product = docx_text(self.client.get("/api/v1/protocol-forms/product.docx").content)
        for fragment in ("по программированию продуктовому", "Судья 3", "ОБЩЕЕ", "Потенциал (10%)", "М.П."):
            self.assertIn(fragment, product)
        self.assertEqual(self.client.get("/api/v1/protocol-forms/other.docx").status_code, 422)


if __name__ == "__main__":
    unittest.main()
